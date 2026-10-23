"""Turn files into Records. One function per format; `load_path` picks by extension.

Rules shared by every loader:
- text is normalised (Unicode NFKC, collapsed spaces), headings become "## heading"
  lines so the structure-aware chunker can see them;
- table-like files (CSV, JSON arrays, JSONL) produce one Record per row, because
  each row is an independent fact with its own id;
- metadata found in the text (dates, ids such as INC-2026-016 or CHG-3340) is
  extracted once here, so search can filter and boost on it later.
"""
import csv
import io
import json
import re
import unicodedata
from pathlib import Path

from .records import Record

ID_RE = re.compile(r"\b[A-Z]{2,5}(?:-\d{2,5})+\b")      # INC-2026-016, CHG-3340, TCK-2026-03-0012
DATE_RE = re.compile(r"\b(20\d\d)-(0[1-9]|1[0-2])-([0-2]\d|3[01])\b")


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def text_meta(text: str) -> dict:
    """Ids and dates mentioned in the text. Dates are kept as YYYY-MM-DD strings."""
    dates = sorted({"-".join(m) for m in DATE_RE.findall(text)})
    return {"ids": sorted(set(ID_RE.findall(text))), "dates": dates}


def _record(doc_id, source, fmt, title, text, **meta):
    text = normalise(text)
    return Record(doc_id, source, fmt, title.strip() or doc_id, text, {**text_meta(text), **meta})


# ------------------------------------------------------------------ per format
def load_pdf(data: bytes, source: str):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [p.extract_text() or "" for p in reader.pages]
    title = (reader.metadata.title if reader.metadata and reader.metadata.title else "") or Path(source).stem
    return [_record(source, source, "pdf", title, "\n\n".join(pages), pages=len(pages))]


def load_docx(data: bytes, source: str):
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            lines.append("\n## " + p.text.strip())
        elif "list" in style:
            lines.append("- " + p.text.strip())
        else:
            lines.append(p.text.strip())
    for table in d.tables:                    # tables become "cell | cell | cell" lines
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    title = d.core_properties.title or Path(source).stem
    return [_record(source, source, "docx", title, "\n".join(lines))]


def load_html(data: bytes, source: str):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    meta = {m.get("name"): m.get("content") for m in soup.find_all("meta") if m.get("name")}
    for h in soup.find_all(["h1", "h2", "h3"]):
        h.string = "## " + h.get_text(" ", strip=True)
    title = soup.title.get_text(" ", strip=True) if soup.title else Path(source).stem
    if soup.title:
        soup.title.decompose()
    text = soup.get_text("\n", strip=True)
    return [_record(source, source, "html", title, text, html_meta=meta)]


def load_md(data: bytes, source: str):
    text = data.decode("utf-8", errors="replace")
    m = re.search(r"^#\s+(.+)$", text, re.M)
    text = re.sub(r"^#{1,6}\s+", "## ", text, flags=re.M)     # one heading marker for every level
    return [_record(source, source, "md", m.group(1) if m else Path(source).stem, text)]


def _row_text(row: dict) -> str:
    return "\n".join(f"{k}: {v}" for k, v in row.items() if v not in (None, ""))


def load_csv(data: bytes, source: str):
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8", errors="replace"))))
    out = []
    for n, row in enumerate(rows):
        row_id = next(iter(row.values())) or str(n)          # first column is the row id
        out.append(_record(f"{source}#{row_id}", source, "csv", f"{Path(source).stem} {row_id}",
                           _row_text(row), row=row))
    return out


def load_json(data: bytes, source: str):
    obj = json.loads(data)
    items = obj if isinstance(obj, list) else [obj]
    out = []
    for n, item in enumerate(items):
        if not isinstance(item, dict):
            item = {"value": item}
        row_id = next((str(v) for k, v in item.items() if k.endswith("_id") or k == "id"), str(n))
        out.append(_record(f"{source}#{row_id}", source, "json", f"{Path(source).stem} {row_id}",
                           _row_text(item), row=item))
    return out


def load_jsonl(data: bytes, source: str):
    """JSONL records that already carry doc_id / title / text (the HotpotQA corpus)."""
    out = []
    for line in data.decode("utf-8").splitlines():
        if line.strip():
            o = json.loads(line)
            out.append(_record(o["doc_id"], o.get("source", source), "jsonl", o.get("title", ""), o["text"]))
    return out


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".html": load_html, ".htm": load_html, ".md": load_md,
           ".markdown": load_md, ".csv": load_csv, ".json": load_json, ".jsonl": load_jsonl, ".txt": load_md}


class UnsupportedFormat(ValueError):
    pass


def load_bytes(data: bytes, source: str):
    ext = Path(source).suffix.lower()
    if ext not in LOADERS:
        raise UnsupportedFormat(f"{ext or 'no extension'} is not supported; use one of {sorted(LOADERS)}")
    return LOADERS[ext](data, source)


def load_path(path: Path, root: Path):
    return load_bytes(path.read_bytes(), path.relative_to(root).as_posix())
